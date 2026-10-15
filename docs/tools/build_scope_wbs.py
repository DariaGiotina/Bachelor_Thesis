"""Build the scope statement and the WBS of the thesis project (Software Project Management course).

Outputs (all regenerated on every run, so edit this script rather than the files):
  docs/Enuntul_domeniului_Giotina.docx   scope statement, the 10 sections of the course + strategy and effort
  docs/WBS_Giotina.docx                  coded WBS (figure + indented table), WBS dictionary, activities of 1.6.2
  docs/diagrams/wbs.puml                 PlantUML source of the WBS tree
  docs/diagrams/wbs.png, wbs.svg         rendered with --render (sends the .puml text to the plantuml.com server)

The WBS data (codes, names, hours, status) is defined once in WBS below and feeds every output.
Usage: python docs/tools/build_scope_wbs.py [--render]
"""
import sys
import urllib.request
import zlib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parents[1]
DIAG = DOCS / "diagrams"
DATE = "15.10.2026"
VERSION = "1.0"
PROJECT = "Companion AI pentru îngrijirea pielii"
PROJECT_LONG = ("Companion AI pentru îngrijirea pielii – aplicație mobilă cu analiza pielii, informații despre "
                "ingrediente și rutine personalizate (lucrare de licență și articol de cercetare)")

# Status on DATE: D = realizat, P = în lucru, N = planificat
STATUS = {"D": "realizat", "P": "în lucru", "N": "planificat"}

# (code, name, hours, status); level-2 rows have hours None (sum of children)
WBS = [
    ("1.1", "Managementul proiectului", None, None),
    ("1.1.1", "Carta proiectului", 8, "D"),
    ("1.1.2", "Enunțul domeniului", 8, "D"),
    ("1.1.3", "Planul proiectului (WBS, dicționar, calendar)", 12, "P"),
    ("1.1.4", "Rapoartele de stare (jurnalul lunar)", 12, "P"),
    ("1.2", "Cercetare și cerințe", None, None),
    ("1.2.1", "Studiul bibliografic", 40, "P"),
    ("1.2.2", "Întrebările de cercetare", 16, "P"),
    ("1.2.3", "Specificația cerințelor aplicației", 24, "N"),
    ("1.2.4", "Registrul surselor de date și al licențelor", 12, "P"),
    ("1.3", "Proiectare", None, None),
    ("1.3.1", "Arhitectura sistemului", 20, "P"),
    ("1.3.2", "Modelul de date (EER, schema)", 20, "D"),
    ("1.3.3", "Prototipul interfeței (Figma)", 32, "P"),
    ("1.3.4", "Specificația API", 12, "N"),
    ("1.4", "Componenta ML de analiză a pielii", None, None),
    ("1.4.1", "Setul de date pregătit (audit, etichete, împărțiri)", 40, "D"),
    ("1.4.2", "Modelul de referință pe imagine (E1)", 40, "D"),
    ("1.4.3", "Modelele multimodale imagine + chestionar (E1–E2)", 60, "D"),
    ("1.4.4", "Raportul pe tonuri de piele și calibrare (E3–E4)", 40, "P"),
    ("1.4.5", "Comparația cu modelele fundație și transferul extern (E5–E6)", 40, "P"),
    ("1.4.6", "Modulul de explicabilitate Grad-CAM", 16, "P"),
    ("1.4.7", "Serviciul de inferență ML", 24, "N"),
    ("1.5", "Backend (API FastAPI)", None, None),
    ("1.5.1", "Modulul de autentificare și conturi", 32, "N"),
    ("1.5.2", "API-ul de analiză a pielii", 24, "N"),
    ("1.5.3", "API-ul de produse și ingrediente", 32, "N"),
    ("1.5.4", "API-ul de rutine, jurnal și progres", 32, "N"),
    ("1.5.5", "Stocarea fotografiilor și drepturile GDPR", 24, "N"),
    ("1.6", "Aplicația mobilă (React Native + Expo)", None, None),
    ("1.6.1", "Ecranele de cont și de pornire", 24, "N"),
    ("1.6.2", "Fluxul de analiză a pielii", 48, "N"),
    ("1.6.3", "Scannerul de coduri de bare și fișa ingredientelor", 40, "N"),
    ("1.6.4", "Constructorul de rutine", 36, "N"),
    ("1.6.5", "Jurnalul, pozele de progres și mementourile", 40, "N"),
    ("1.7", "Testare", None, None),
    ("1.7.1", "Suita de teste de integrare (API + ML)", 24, "N"),
    ("1.7.2", "Testele end-to-end ale aplicației", 24, "N"),
    ("1.7.3", "Raportul de securitate și GDPR", 20, "N"),
    ("1.7.4", "Raportul testului beta", 24, "N"),
    ("1.8", "Articolul de cercetare", None, None),
    ("1.8.1", "Manuscrisul complet", 60, "P"),
    ("1.8.2", "Figurile și tabelele rezultatelor", 20, "P"),
    ("1.8.3", "Versiunea trimisă revistei", 24, "N"),
    ("1.9", "Teza de licență", None, None),
    ("1.9.1", "Capitolele introductive și teoretice (1–3)", 60, "P"),
    ("1.9.2", "Capitolele de proiectare și implementare (4–7)", 70, "P"),
    ("1.9.3", "Capitolele de rezultate și concluzii (8–9)", 40, "P"),
    ("1.9.4", "Versiunea finală formatată", 16, "N"),
    ("1.9.5", "Prezentarea pentru susținere", 20, "N"),
    ("1.10", "Livrarea aplicației", None, None),
    ("1.10.1", "Aplicația instalată (build + API în UE)", 20, "N"),
    ("1.10.2", "Codul sursă publicat și README", 10, "N"),
    ("1.10.3", "Manualul utilizatorului", 16, "N"),
]


def level(code: str) -> int:
    return code.count(".") + 1


def children(parent: str):
    return [r for r in WBS if r[0].startswith(parent + ".") and level(r[0]) == level(parent) + 1]


def hours(code: str) -> int:
    row = next(r for r in WBS if r[0] == code)
    return row[2] if row[2] is not None else sum(hours(c[0]) for c in children(code))


TOTAL = sum(hours(r[0]) for r in WBS if level(r[0]) == 2)
L2 = [r for r in WBS if level(r[0]) == 2]
WPS = [r for r in WBS if r[2] is not None]


# ---------------------------------------------------------------- PlantUML
def puml() -> str:
    lines = [
        "@startwbs",
        "scale 1.45",
        "<style>",
        "wbsDiagram {",
        "  FontName Arial",
        "  node { FontSize 12; Padding 7; Margin 5; RoundCorner 8; LineColor #44546A; MaximumWidth 165 }",
        "  arrow { LineColor #44546A }",
        "  .root { BackgroundColor #1F3864; FontColor #FFFFFF; FontSize 15; FontStyle bold; MaximumWidth 420 }",
        "  .phase { BackgroundColor #D6E4F0; FontStyle bold; FontSize 12 }",
        "  .done { BackgroundColor #C6EFCE }",
        "  .prog { BackgroundColor #FFEB9C }",
        "  .plan { BackgroundColor #FFFFFF }",
        "}",
        "</style>",
        f"* 1 {PROJECT}\\n{TOTAL} h <<root>>",
    ]
    cls = {"D": "done", "P": "prog", "N": "plan"}
    for code, name, h, st in WBS:
        stars = "*" * level(code)
        if h is None:
            lines.append(f"{stars} {code} {name}\\n{hours(code)} h <<phase>>")
        else:
            lines.append(f"{stars} {code} {name}\\n{h} h <<{cls[st]}>>")
    lines += [
        "legend right",
        "  <back:#C6EFCE>   </back> realizat   <back:#FFEB9C>   </back> în lucru   "
        f"<back:#FFFFFF>   </back> planificat   (stare la {DATE}; h = ore de efort estimate)",
        "endlegend",
        "@endwbs",
    ]
    return "\n".join(lines) + "\n"


def _encode(text: str) -> str:
    """PlantUML text encoding: raw deflate + PlantUML's base64 alphabet."""
    data = zlib.compress(text.encode("utf-8"), 9)[2:-4]
    alpha = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz-_"
    out = []
    for i in range(0, len(data), 3):
        b = data[i:i + 3] + bytes(3 - len(data[i:i + 3]))
        n = (b[0] << 16) | (b[1] << 8) | b[2]
        out.append("".join(alpha[(n >> s) & 63] for s in (18, 12, 6, 0)))
    return "".join(out)


def render(text: str) -> None:
    enc = _encode(text)
    for fmt in ("png", "svg"):
        req = urllib.request.Request(f"https://www.plantuml.com/plantuml/{fmt}/{enc}",
                                     headers={"User-Agent": "build_scope_wbs"})
        with urllib.request.urlopen(req, timeout=60) as r:
            (DIAG / f"wbs.{fmt}").write_bytes(r.read())


# ---------------------------------------------------------------- docx helpers
def new_doc() -> Document:
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.paragraph_format.space_after = Pt(4)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    return d


def para(d, text, bold_prefix=None, italic=False, size=None, align=None):
    p = d.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def bullets(d, items):
    for it in items:
        if isinstance(it, tuple):
            p = d.add_paragraph(style="List Bullet")
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            d.add_paragraph(it, style="List Bullet")


def shade(cell, hex_fill):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc.append(shd)


def table(d, rows, widths=None, header=True, size=10, first_col_bold=False):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = ""
            run = c.paragraphs[0].add_run(str(v))
            run.font.size = Pt(size)
            if (header and i == 0) or (first_col_bold and j == 0):
                run.bold = True
            if header and i == 0:
                shade(c, "D6E4F0")
            if widths:
                c.width = Cm(widths[j])
    d.add_paragraph()
    return t


def title_block(d, title, subtitle):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(20), RGBColor(0x1F, 0x38, 0x64)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(12)
    table(d, [
        ["Proiect", PROJECT_LONG],
        ["Disciplina", "Managementul proiectelor software (anul IV, semestrul 1)"],
        ["Instituție", "Universitatea din Craiova, Facultatea de Automatică, Calculatoare și Electronică"],
        ["Autor / manager de proiect", "Daria Giotina"],
        ["Profesor coordonator (sponsor)", "Stefan Popescu"],
        ["Versiune / dată", f"{VERSION} / {DATE}"],
    ], widths=[5.0, 11.6], header=False, first_col_bold=True)


# ---------------------------------------------------------------- scope statement
def scope_statement() -> None:
    d = new_doc()
    title_block(d, "Enunțul domeniului proiectului", "Project scope statement")
    para(d, "Documentul descrie ce livrează proiectul și ce nu livrează, pe baza cartei proiectului din 09.10.2026. "
            "Împreună cu WBS-ul și dicționarul WBS (documentul WBS_Giotina.docx) formează baza de referință a "
            "domeniului. După aprobare, orice modificare trece printr-o cerere de schimbare (analiza impactului "
            "asupra timpului, efortului și calității, decizia coordonatorului, actualizarea documentelor).",
         italic=True)

    d.add_heading("1. Scopul proiectului (justificarea)", 1)
    para(d, "Problemele obișnuite ale pielii (acnee, eczemă, roșeață) sunt fotografiate de utilizatori cu telefoane "
            "obișnuite, însă cercetarea s-a concentrat mai ales pe cancerul de piele și pe imagini clinice. "
            "Aplicațiile de skincare pentru consumatori cer de obicei un chestionar, dar utilizatorii reali sar peste "
            "întrebări, iar seturile publice de imagini acoperă inegal tonurile de piele, astfel încât o acuratețe "
            "medie poate ascunde un comportament inechitabil pentru pielea închisă la culoare. În plus, listele de "
            "ingrediente de pe produse sunt greu de înțeles pentru un nespecialist.")
    para(d, "Proiectul răspunde acestor nevoi printr-o aplicație mobilă strict informativă (nu este dispozitiv "
            "medical) și printr-un studiu de cercetare care verifică, pe setul de date public SCIN (Google), dacă un "
            "chestionar scurt ajută clasificarea față de fotografia singură, ce se întâmplă când lipsesc răspunsuri "
            "și dacă efectele sunt aceleași pe toate tonurile de piele.")

    d.add_heading("2. Obiectivele (SMART)", 1)
    para(d, "SMART = specific, măsurabil, realizabil, relevant și încadrat în timp. Lunile M1–M10 sunt lunile "
            "proiectului: M1 = august 2026, …, M10 = mai 2027.")
    table(d, [
        ["Cod", "Obiectiv", "Termen"],
        ["O1", "Rezultatele experimentelor E1–E4 pe SCIN (10 seed-uri, intervale de încredere bootstrap de 95%, "
               "rezultate pe grupe de ton de piele) sunt reproductibile din scripturile din depozit.", "30.11.2026 (M4)"],
        ["O2", "Articolul de cercetare (EN), aprobat de coordonator, este trimis la revista aleasă împreună cu "
               "acesta.", "31.12.2026 (M5)"],
        ["O3", "MVP-ul rulează pe un telefon Android real: analiza pielii (foto + chestionar) cu rezultat în cel "
               "mult 10 s, scanarea codului de bare, rutine, jurnal, poze de progres și mementouri.", "31.03.2027 (M8)"],
        ["O4", "Testul beta cu cel puțin 5 utilizatori este încheiat, cu zero defecte critice deschise.",
         "30.04.2027 (M9)"],
        ["O5", "Teza de licență (RO), aprobată de coordonator, este predată.", "31.05.2027 (M10)"],
    ], widths=[1.2, 11.8, 3.6])

    d.add_heading("3. Descrierea produsului și cerințele de nivel înalt", 1)
    para(d, "Produsul are trei părți: aplicația mobilă (React Native + Expo), API-ul backend (FastAPI, interfața "
            "prin care aplicația citește și scrie date) cu baza de date MongoDB Atlas în UE, și serviciul ML de "
            "inferență, care primește fotografia și răspunsurile și întoarce categoria cea mai probabilă dintr-un set "
            "mic de clase (acnee, eczemă/dermatită, tip roșeață, altele), un scor de încredere, harta Grad-CAM "
            "(zonele imaginii care au contat cel mai mult pentru decizie) și ieșirea explicită „incert – consultați "
            "un specialist” când încrederea este scăzută.")
    para(d, "Cerințe de nivel înalt:")
    bullets(d, [
        "Contul utilizatorului: înregistrare, autentificare (parole Argon2id, tokenuri JWT cu reîmprospătare "
        "rotativă), consimțământ explicit pentru fotografii și date de sănătate.",
        "Analiza pielii din fotografie + chestionar scurt; fiecare întrebare se poate sări, iar răspunsul lipsă se "
        "înregistrează explicit ca lipsă.",
        "Scanarea codului de bare al unui produs și afișarea ingredientelor cu funcția și marcajele lor de risc "
        "(surse: Open Beauty Facts, CosIng, CIR).",
        "Constructorul de rutine (dimineață/seară), jurnalul zilnic, pozele de progres și mementourile.",
        "Drepturile GDPR din aplicație: exportul datelor și ștergerea contului împreună cu toate fotografiile.",
        "Niciun text de diagnostic sau tratament; fiecare rezultat afișează scorul de încredere și avertismentul "
        "că aplicația are rol informativ.",
    ])

    d.add_heading("4. Livrabilele", 1)
    para(d, "Livrabilele de mai jos devin ramurile de nivel 2 ale WBS-ului (codul WBS este în paranteză).")
    table(d, [
        ["Cod", "Livrabil", "Tip"],
        ["L1", "Aplicația mobilă instalată și codul ei sursă (1.6, 1.10)", "final"],
        ["L2", "API-ul backend, cu stocarea fotografiilor în UE (1.5)", "final"],
        ["L3", "Componenta ML: setul de date pregătit, modelele, serviciul de inferență, rapoartele E1–E6 (1.4)",
         "final"],
        ["L4", "Articolul de cercetare în limba engleză (1.8)", "final"],
        ["L5", "Teza de licență în limba română și prezentarea pentru susținere (1.9)", "final"],
        ["L6", "Documentele de cercetare și proiectare: studiul bibliografic, întrebările de cercetare, "
               "cerințele, arhitectura, modelul de date, prototipul Figma, specificația API (1.2, 1.3)",
         "intermediar"],
        ["L7", "Rapoartele de testare: integrare, end-to-end, securitate și GDPR, beta (1.7)", "intermediar"],
        ["L8", "Manualul utilizatorului (1.10.3)", "documentare"],
        ["L9", "Documentele de management: carta, enunțul domeniului, WBS-ul, rapoartele de stare (1.1)",
         "management"],
    ], widths=[1.2, 12.4, 3.0])

    d.add_heading("5. Domeniul inclus", 1)
    bullets(d, [
        "Tot lucrul din WBS: managementul proiectului, cercetarea, proiectarea, componenta ML, backend-ul, "
        "aplicația mobilă, testarea, articolul, teza și livrarea aplicației.",
        "Experimentele E1–E4 (obligatorii) și E5–E6 (comparația cu modelele fundație și transferul pe alt set de "
        "date, dacă datele externe se pot obține legal).",
        "Testul beta cu voluntari (colegi, prieteni) pe telefoane Android.",
    ])

    d.add_heading("6. Domeniul exclus", 1)
    bullets(d, [
        "Diagnosticul sau tratamentul medical și orice afirmație de acuratețe de nivel medical.",
        "Clasamentele de tendințe și popularitate, monetizarea, linkurile afiliate și marketplace-ul în aplicație.",
        "Teleconsultația cu un dermatolog.",
        "Publicarea în Google Play sau App Store (beta-ul se distribuie ca build de test).",
        "Mentenanța și suportul pentru utilizatori după 31.05.2027.",
        "Antrenarea modelelor pe Fitzpatrick17k (imaginile se folosesc doar la testare și se șterg după cercetare).",
        "Funcțiile „stretch” (harta facială pe mai multe zone, inferența pe telefon, partajarea socială, "
        "avertizările de conflict între ingrediente): intră doar printr-o cerere de schimbare aprobată, dacă "
        "rămâne timp după M8.",
    ])
    para(d, "Articolul acoperă exclusiv partea de ML; aplicația apare doar în teză.", italic=True)

    d.add_heading("7. Criteriile de acceptare", 1)
    para(d, "Macro-F1 este media scorului F1 (media armonică dintre precizie și recall) calculată separat pentru "
            "fiecare clasă, astfel încât clasele rare cântăresc la fel de mult ca cele frecvente.")
    table(d, [
        ["Livrabil", "Criteriu de acceptare verificabil"],
        ["L1 Aplicația", "Testele end-to-end ale fluxurilor principale (cont, analiză, scanare, rutină, jurnal) trec; "
                         "rezultatul analizei apare în cel mult 10 s (mediana a 20 de încercări pe Wi-Fi); niciun "
                         "termen din lista de termeni de diagnostic interziși nu apare în interfață; zero defecte "
                         "critice deschise la predare."],
        ["L2 API-ul", "Testele de integrare trec; un utilizator nu poate citi datele altui utilizator (test de "
                      "izolare); exportul și ștergerea contului elimină toate documentele și fotografiile "
                      "utilizatorului (verificat în baza de date și în stocare)."],
        ["L3 Componenta ML", "Tabelele E1–E4 se regenerează cu câte o comandă din README; fiecare rezultat raportează "
                             "media ± abaterea pe 10 seed-uri, macro-F1, rezultatele pe clase și pe grupe de ton "
                             "(Fitzpatrick, Monk) și ECE (eroarea de calibrare); testele pytest trec."],
        ["L4 Articolul", "Manuscrisul este aprobat în scris de coordonator și confirmarea de primire a revistei "
                         "există până la 31.12.2026."],
        ["L5 Teza", "Respectă șablonul facultății, trece verificarea antiplagiat sub pragul facultății și este "
                    "aprobată de coordonator până la 31.05.2027."],
        ["L7 Testarea", "Raportul beta conține răspunsurile a cel puțin 5 utilizatori și lista defectelor găsite, "
                        "fiecare cu starea ei."],
        ["L8 Manualul", "Acoperă toate fluxurile din cerințele de nivel înalt, cu capturi de ecran."],
    ], widths=[3.4, 13.2])

    d.add_heading("8. Constrângerile", 1)
    bullets(d, [
        "Termene fixe: articolul în decembrie 2026, teza la 31.05.2027; 10 luni (01.08.2026–31.05.2027).",
        "Un singur membru în echipă pentru ML, backend, frontend, testare și documentare "
        f"({TOTAL} de ore estimate în WBS, aproximativ {round(TOTAL / 43)} de ore pe săptămână timp de 43 de "
        "săptămâni).",
        "Fără buget extern; calcul pe un PC cu NVIDIA RTX 3060, cu Google Colab sau Kaggle ca rezervă.",
        "Licența SCIN: folosire pentru cercetare, fără redistribuirea imaginilor, re-identificarea persoanelor "
        "interzisă; imaginile Fitzpatrick17k doar pentru testare, șterse după cercetare.",
        "GDPR: fotografiile faciale și datele despre piele sunt categorii speciale de date; stocare doar în UE, "
        "cu consimțământ explicit.",
        "Tehnologii stabilite: PyTorch, FastAPI, MongoDB Atlas (UE), React Native + Expo.",
    ])

    d.add_heading("9. Ipotezele", 1)
    para(d, "Fiecare ipoteză este verificată pe parcurs; dacă se dovedește falsă, devine risc și intră în registrul "
            "riscurilor din jurnalul proiectului.")
    bullets(d, [
        "SCIN rămâne disponibil sub aceeași licență pe toată durata proiectului.",
        "Nivelurile gratuite ale MongoDB Atlas și ale stocării de obiecte compatibile S3 din UE ajung pentru MVP "
        "și pentru beta.",
        "Open Beauty Facts conține suficiente produse cosmetice cu listă de ingrediente pentru demonstrarea "
        "scannerului.",
        "Cel puțin 5 colegi sau prieteni acceptă să testeze aplicația în M9.",
        "Coordonatorul revizuiește manuscrisul și capitolele tezei în cel mult o săptămână.",
        "Placa RTX 3060 (12 GB) ajunge pentru antrenarea modelelor EfficientNet și MobileNetV3.",
    ])

    d.add_heading("10. Aprobările și responsabilitățile", 1)
    table(d, [
        ["Rol", "Nume", "Responsabilitate", "Semnătura / data"],
        ["Sponsor", "Stefan Popescu, profesor coordonator",
         "Aprobă enunțul domeniului și cererile de schimbare; revizuiește articolul și teza.", ""],
        ["Client", "Comisia de licență a facultății",
         "Evaluează teza și aplicația la susținere (acceptarea finală).", "–"],
        ["Manager de proiect", "Daria Giotina",
         "Răspunde de realizarea tuturor livrabilelor și execută tot lucrul din WBS.", ""],
    ], widths=[3.0, 4.2, 6.6, 2.8])

    d.add_heading("Anexa 1. Strategia", 1)
    para(d, "Proiectul se face pe faze lunare (M1–M10), cu partea de ML înaintea aplicației, pentru că articolul are "
            "termenul cel mai apropiat. Fiecare lună se încheie cu o intrare în jurnalul proiectului (raportul de "
            "stare) și cu un commit în depozitul GitHub. Funcțiile aplicației se dezvoltă incremental: întâi fluxul "
            "de analiză, apoi scannerul și rutinele, apoi urmărirea progresului, ca un MVP utilizabil să existe cât "
            "mai devreme. Întrebările de cercetare și domeniul articolului se stabilesc împreună cu coordonatorul.")

    d.add_heading("Anexa 2. Estimarea efortului și a costurilor", 1)
    rows = [["Cod WBS", "Livrabil major", "Efort (ore)"]]
    rows += [[c, n, str(hours(c))] for c, n, _, _ in L2]
    rows.append(["", "Total", str(TOTAL)])
    table(d, rows, widths=[2.0, 11.6, 3.0])
    para(d, "Costuri directe: 0 lei pentru hardware și software (echipament propriu, instrumente gratuite, date "
            "publice). Găzduirea folosește niveluri gratuite; costul lunar real se confirmă la configurarea "
            "serviciilor din M5. Nu se face o analiză cost-beneficiu financiară: beneficiul proiectului este "
            "academic (teza, articolul) și practic (un instrument informativ pentru utilizatori).")

    d.save(DOCS / "Enuntul_domeniului_Giotina.docx")


# ---------------------------------------------------------------- WBS document
DICTIONARY = [
    ("1.4.3", "Modelele multimodale imagine + chestionar (E1–E2)", [
        ("Descrierea lucrului", "Modele care combină caracteristicile imaginii (EfficientNet-B0 preantrenat) cu "
                                "răspunsurile codificate ale chestionarului și masca de răspunsuri lipsă: fuziune "
                                "tardivă prin concatenare, fuziune cu poartă (gated) și FiLM, plus antrenarea cu "
                                "răspunsuri ascunse aleator (answer dropout) pentru experimentul E2."),
        ("Criterii de acceptare", "Fiecare variantă este antrenată pe 10 seed-uri; tabelele E1 și E2 conțin macro-F1 "
                                  "pe test cu intervale de încredere bootstrap și comparația pereche cu modelul doar "
                                  "pe imagine; testele unitare ale modelului trec."),
        ("Responsabil", "Daria Giotina execută (R) și răspunde (A); coordonatorul este consultat asupra "
                        "interpretării (C)."),
        ("Efort estimat", "60 de ore, în M3–M4"),
        ("Dependențe", "1.4.1 Setul de date pregătit; 1.4.2 Modelul de referință pe imagine"),
        ("Ipoteze și riscuri", "Risc: chestionarul nu aduce câștig măsurabil (s-a confirmat; se tratează ca rezultat "
                               "și se explică prin controale de metodă, fără a forța un rezultat pozitiv)."),
    ]),
    ("1.4.4", "Raportul pe tonuri de piele și calibrare (E3–E4)", [
        ("Descrierea lucrului", "Raport care împarte rezultatele pe grupe de ton de piele (Fitzpatrick I–II, III–IV, "
                                "V–VI; Monk 1–3, 4–6, 7–10), cu performanța grupei celei mai slabe, calibrarea (ECE) "
                                "și curbele risc-acoperire pentru ieșirea „incert”."),
        ("Criterii de acceptare", "Fiecare grupă are macro-F1 cu interval de încredere de 95% și numărul de cazuri; "
                                  "testele de diferență între grupe sunt corectate Holm; grupele cu sub 30 de cazuri "
                                  "sunt marcate ca descriptive."),
        ("Responsabil", "Daria Giotina (R, A); coordonatorul (C)."),
        ("Efort estimat", "40 de ore, în M4"),
        ("Dependențe", "1.4.3 Modelele multimodale; etichetele de ton eFST și eMST din 1.4.1"),
        ("Ipoteze și riscuri", "Puține cazuri pentru tonurile închise (FST VI: 68 de cazuri) dau intervale largi; "
                               "limita se raportează deschis."),
    ]),
    ("1.5.1", "Modulul de autentificare și conturi", [
        ("Descrierea lucrului", "Endpoint-urile FastAPI pentru înregistrare, autentificare, reîmprospătarea și "
                                "revocarea sesiunii și consimțământul GDPR. Parolele sunt păstrate ca hash Argon2id; "
                                "accesul folosește tokenuri JWT de scurtă durată și tokenuri de reîmprospătare "
                                "rotative, stocate ca hash."),
        ("Criterii de acceptare", "Testele unitare și de integrare ale modulului trec; un token de reîmprospătare "
                                  "folosit a doua oară este respins și revocă sesiunea; nicio parolă sau token nu "
                                  "apare în clar în baza de date sau în loguri."),
        ("Responsabil", "Daria Giotina (R, A)."),
        ("Efort estimat", "32 de ore, în M5"),
        ("Dependențe", "1.3.2 Modelul de date; 1.3.4 Specificația API"),
        ("Ipoteze și riscuri", "Ipoteză: nivelul gratuit MongoDB Atlas din UE este suficient. Risc: o greșeală de "
                               "securitate în implementarea proprie; se reduce prin testele din 1.7.3."),
    ]),
    ("1.6.2", "Fluxul de analiză a pielii", [
        ("Descrierea lucrului", "Ecranele aplicației prin care utilizatorul face o fotografie ghidată, completează "
                                "chestionarul (cu opțiunea de a sări peste orice întrebare) și vede rezultatul: "
                                "categoria, scorul de încredere, harta Grad-CAM și mesajul „consultați un specialist” "
                                "când rezultatul este incert."),
        ("Criterii de acceptare", "Testele componentelor trec; pe un telefon Android rezultatul apare în cel mult 10 s; "
                                  "o întrebare sărită ajunge la API marcată ca lipsă; ecranul nu conține termeni de "
                                  "diagnostic."),
        ("Responsabil", "Daria Giotina (R, A); utilizatorii beta sunt informați și dau feedback (I)."),
        ("Efort estimat", "48 de ore, în M6 (activitățile de mai jos)"),
        ("Dependențe", "1.3.3 Prototipul interfeței; 1.5.2 API-ul de analiză; 1.4.7 Serviciul de inferență"),
        ("Ipoteze și riscuri", "Ipoteză: camera telefonului de test oferă calitate suficientă. Risc: timpul de răspuns "
                               "al inferenței în cloud depășește 10 s; măsură: imagine redimensionată înainte de "
                               "încărcare."),
    ]),
]

ACTIVITIES = [
    ("A1", "Proiectează navigarea și stările ecranelor de analiză", 4),
    ("A2", "Implementează captura foto ghidată (expo-camera)", 10),
    ("A3", "Implementează chestionarul cu opțiunea „sari peste”", 8),
    ("A4", "Implementează încărcarea fotografiei prin URL presemnat și apelul de analiză", 8),
    ("A5", "Implementează ecranul de rezultat (categorie, încredere, Grad-CAM, avertisment)", 10),
    ("A6", "Scrie testele componentelor", 6),
    ("A7", "Code review și corecturi", 2),
]


def wbs_document() -> None:
    d = new_doc()
    title_block(d, "Structura de descompunere a lucrărilor (WBS)", "Work Breakdown Structure")
    para(d, "WBS-ul descompune ierarhic tot lucrul proiectului în livrabile (substantive, nu activități). Nivelul 1 "
            "este proiectul, nivelul 2 sunt livrabilele majore din enunțul domeniului, iar nivelul 3 sunt pachetele "
            "de lucru: unitatea care se estimează, se atribuie și se verifică. Codul fiecărui element începe cu "
            "codul părintelui. Pachetele respectă regula 8/80 (între 8 și 80 de ore de efort).")
    para(d, "Diagrama a fost realizată în PlantUML (sintaxa @startwbs), iar sursa editabilă este "
            "docs/diagrams/wbs.puml; ea se regenerează cu docs/tools/build_scope_wbs.py.", italic=True)

    # the tree is wide, so it gets its own landscape page
    portrait = d.sections[0]
    land = d.add_section(WD_SECTION.NEW_PAGE)
    land.orientation = WD_ORIENT.LANDSCAPE
    land.page_width, land.page_height = portrait.page_height, portrait.page_width
    land.left_margin = land.right_margin = land.top_margin = land.bottom_margin = Cm(1.2)
    d.add_heading("1. Diagrama WBS", 1)
    png = DIAG / "wbs.png"
    if png.exists():
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(png), width=land.page_width - land.left_margin - land.right_margin)
        para(d, f"Figura 1. WBS-ul proiectului; culoarea pachetelor arată starea la {DATE}.", italic=True,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    back = d.add_section(WD_SECTION.NEW_PAGE)
    back.orientation = WD_ORIENT.PORTRAIT
    back.page_width, back.page_height = portrait.page_width, portrait.page_height
    back.left_margin, back.right_margin = portrait.left_margin, portrait.right_margin
    back.top_margin, back.bottom_margin = portrait.top_margin, portrait.bottom_margin

    d.add_heading("2. WBS-ul ca tabel (listă indentată)", 1)
    rows = [["Cod", "Element", "Efort (ore)", "Stare"]]
    rows.append(["1", PROJECT, str(TOTAL), ""])
    for code, name, h, st in WBS:
        indent = "    " * (level(code) - 2)
        rows.append([code, indent + name, str(hours(code)), STATUS[st] if st else ""])
    t = table(d, rows, widths=[1.6, 10.4, 2.2, 2.4], size=9)
    for i, r in enumerate(rows[1:], start=1):
        if i == 1 or level(r[0]) <= 2:
            for c in t.rows[i].cells:
                shade(c, "EAF1F8")
                for run in c.paragraphs[0].runs:
                    run.bold = True

    d.add_heading("3. Verificarea WBS-ului", 1)
    bullets(d, [
        ("Regula 100%: ", "ramurile acoperă tot lucrul din enunțul domeniului, inclusiv managementul proiectului "
                          "(1.1); fiecare livrabil L1–L9 are ramura lui, iar efortul nivelului 2 este suma "
                          "pachetelor de lucru."),
        ("Fără dubluri: ", "testele unitare fac parte din pachetul componentei testate, iar ramura 1.7 conține doar "
                           "testele de integrare, end-to-end, securitate și beta; evaluarea modelului este în 1.4, "
                           "iar figurile articolului (1.8.2) doar o formatează pentru revistă."),
        ("Substantive, nu verbe: ", "toate elementele sunt rezultate („Modulul de autentificare”), iar verbele apar "
                                    "doar la activitățile din secțiunea 5."),
        ("Mărimea pachetelor: ", f"{len(WPS)} pachete de lucru, între {min(r[2] for r in WPS)} și "
                                 f"{max(r[2] for r in WPS)} de ore; total {TOTAL} de ore."),
    ])

    d.add_heading("4. Dicționarul WBS", 1)
    para(d, "Fișele pentru patru pachete de lucru. Responsabilitățile urmează matricea RACI: R = execută, "
            "A = răspunde de rezultat, C = consultat, I = informat.")
    for code, name, fields in DICTIONARY:
        d.add_heading(f"Fișa pachetului {code} · {name}", 2)
        table(d, [[k, v] for k, v in fields], widths=[4.0, 12.6], header=False, first_col_bold=True)

    d.add_heading("5. Activitățile pachetului 1.6.2 Fluxul de analiză a pielii", 1)
    rows = [["Cod", "Activitatea", "Efort"]]
    rows += [[c, n, f"{h} h"] for c, n, h in ACTIVITIES]
    total = sum(h for _, _, h in ACTIVITIES)
    rows.append(["", "Total", f"{total} h"])
    table(d, rows, widths=[1.4, 12.6, 2.6])
    assert total == hours("1.6.2"), "activities must add up to the work package effort"
    para(d, f"Suma activităților ({total} h) este egală cu efortul pachetului ({hours('1.6.2')} h). "
            "Ordinea activităților: A1 → (A2 ∥ A3) → A4 → A5 → A6 → A7; captura foto și chestionarul se pot face în "
            "paralel, după care urmează legarea lor de API și ecranul de rezultat. Ordinea și duratele trec apoi în "
            "calendarul proiectului (diagrama Gantt).")

    d.save(DOCS / "WBS_Giotina.docx")


def main() -> None:
    text = puml()
    (DIAG / "wbs.puml").write_text(text, encoding="utf-8")
    if "--render" in sys.argv:
        render(text)
    scope_statement()
    wbs_document()
    print(f"WBS: {len(WPS)} work packages, {TOTAL} h total")


if __name__ == "__main__":
    main()
