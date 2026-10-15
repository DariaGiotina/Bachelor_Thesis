"""Scope statement and WBS of the project (Software Project Management course).

Adds Annex A to the thesis with a summary of the scope statement and the WBS figure. The full documents are
docs/Enuntul_domeniului_Giotina.docx and docs/WBS_Giotina.docx (built by docs/tools/build_scope_wbs.py).
The paper covers only the ML part, so it is not changed.
Usage: python docs/tools/updates/2026-10-15_scope_wbs.py <repo root>
"""
import copy
import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

ROOT = Path(sys.argv[1])
sys.path.insert(0, str(ROOT / "docs" / "tools"))
from docx_tools import DocEditor  # noqa: E402
from build_scope_wbs import L2, TOTAL, WPS, hours  # noqa: E402

PMBOK = ("https://www.pmi.org/pmbok-guide-standards/foundational/pmbok",
         "Project Management Institute, A Guide to the Project Management Body of Knowledge (PMBOK Guide), "
         "ediția a 6-a, 2017")
PLANTUML = ("https://plantuml.com/wbs-diagram", "PlantUML, Work Breakdown Structure (WBS) diagram, documentație")

HEAD = "Anexa A. Planificarea proiectului: enunțul domeniului și WBS"
NEXT = "Bibliografie"

TEXT_A = [
    "Proiectul a fost planificat după procesele de management al domeniului din ghidul PMBOK [{pmbok}]. Domeniul "
    "proiectului înseamnă tot lucrul necesar pentru a produce livrabilele, și numai acesta. Planificarea are trei "
    "documente: carta proiectului (care autorizează proiectul și fixează obiectivele de nivel înalt), enunțul "
    "domeniului (scope statement) și structura de descompunere a lucrărilor (Work Breakdown Structure, WBS) cu "
    "dicționarul ei. Ultimele două, împreună cu enunțul, formează baza de referință a domeniului: după aprobare, "
    "orice schimbare de domeniu se face doar printr-o cerere de schimbare analizată ca impact asupra timpului, "
    "efortului și calității.",
    "Enunțul domeniului (fișierul docs/Enuntul_domeniului_Giotina.docx) descrie scopul proiectului, obiectivele "
    "SMART (specifice, măsurabile, realizabile, relevante și încadrate în timp), produsul și cerințele de nivel "
    "înalt, nouă livrabile, domeniul inclus și cel exclus, criteriile de acceptare verificabile ale fiecărui "
    "livrabil, constrângerile, ipotezele și aprobările. Excluderile scrise explicit (diagnosticul medical, "
    "monetizarea, teleconsultația, publicarea în magazinele de aplicații, mentenanța după susținere și funcțiile "
    "„stretch”) sunt prima apărare împotriva extinderii necontrolate a domeniului (scope creep). Constrângerile "
    "(limite impuse, de exemplu termenele și licența SCIN) sunt separate de ipoteze (lucruri considerate adevărate "
    "fără a fi dovedite, de exemplu disponibilitatea a cel puțin cinci testeri beta), deoarece o ipoteză falsă "
    "devine un risc.",
    "WBS-ul (Figura A.1, fișierul docs/WBS_Giotina.docx) este descompunerea ierarhică a lucrului în livrabile, nu "
    "în activități. Nivelul 1 este proiectul, nivelul 2 sunt cele {n_l2} livrabile majore (managementul "
    "proiectului, cercetarea și cerințele, proiectarea, componenta ML, backend-ul, aplicația mobilă, testarea, "
    "articolul, teza și livrarea aplicației), iar nivelul 3 sunt cele {n_wp} pachete de lucru. Un pachet de lucru "
    "este elementul cel mai de jos al WBS-ului, pentru care efortul se estimează, se atribuie unei persoane și se "
    "verifică printr-un rezultat clar. Codul fiecărui element începe cu codul părintelui (de exemplu 1.4 → 1.4.3), "
    "iar pachetele respectă regula 8/80 (între 8 și 80 de ore de efort). Efortul total estimat este de {total} de "
    "ore, dintre care {ml} pentru componenta ML. WBS-ul respectă regula 100%: ramurile acoperă tot lucrul din "
    "enunț, inclusiv managementul proiectului, iar fiecare lucru apare o singură dată (de exemplu, testele unitare "
    "aparțin pachetului componentei testate, iar ramura de testare conține doar testele de integrare, "
    "end-to-end, de securitate și beta). Diagrama a fost realizată în PlantUML [{plantuml}], iar culoarea "
    "pachetelor arată starea lor la 15.10.2026.",
    "Dicționarul WBS conține fișele a patru pachete de lucru (1.4.3, 1.4.4, 1.5.1 și 1.6.2), fiecare cu descrierea "
    "lucrului, criteriile de acceptare, responsabilul după matricea RACI (R = execută, A = răspunde, "
    "C = consultat, I = informat), efortul, dependențele și riscurile. Pentru pachetul 1.6.2 (fluxul de analiză a "
    "pielii, 48 de ore) au fost derivate șapte activități, a căror sumă este egală cu efortul pachetului; abia la "
    "acest nivel apar verbele și ordinea lucrului, care trec apoi în calendarul proiectului (diagrama Gantt).",
]


def add_toc_level2(ed, before_text, text, page):
    tpl = next(p for p in ed.doc.paragraphs if p.style.name.lower() == "toc 2")
    target = ed.toc_entries(before_text)[0]
    new = copy.deepcopy(tpl._p)
    target._p.addprevious(new)
    ts = Paragraph(new, target._parent)._p.findall(".//" + qn("w:t"))
    ts[0].text, ts[-1].text = text, page
    for t in ts[1:-1]:
        t.text = "\t" if t.text == "\t" else ""


ed = DocEditor(str(ROOT / "docs" / "Teza_Licenta.docx"))
if ed.has_heading(HEAD):
    print("thesis already updated")
    sys.exit(0)
n = {"pmbok": ed.add_reference(PMBOK[1], PMBOK[0]), "plantuml": ed.add_reference(PLANTUML[1], PLANTUML[0]),
     "n_l2": len(L2), "n_wp": len(WPS), "total": TOTAL, "ml": hours("1.4")}
anchor = ed.heading(NEXT)
ed.heading_before(anchor, HEAD, 2)
for t in TEXT_A[:3]:
    ed.body_paragraph_before(anchor, t.format(**n))
par = ed.body_paragraph_before(anchor, "")
par.paragraph_format.first_line_indent = Cm(0)
par.alignment = WD_ALIGN_PARAGRAPH.CENTER
par.add_run().add_picture(str(ROOT / "docs" / "diagrams" / "wbs.png"), width=Cm(16.0))
cap = anchor.insert_paragraph_before(
    "Figura A.1. Structura de descompunere a lucrărilor (WBS) a proiectului, cu efortul estimat pe fiecare element "
    "(verde = realizat, galben = în lucru, alb = planificat). Versiunea la scară mare este docs/diagrams/wbs.svg.",
    style="Caption")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
ed.body_paragraph_before(anchor, TEXT_A[3].format(**n))
add_toc_level2(ed, NEXT, HEAD, "")
try:
    ed.save()
    print("thesis updated")
except PermissionError:
    print("Teza_Licenta.docx LOCKED (open in Word) - not updated")
