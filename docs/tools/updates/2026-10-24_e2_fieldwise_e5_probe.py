"""E2 completed (field-wise removal, naturally missing questionnaires, curve) and E5 (frozen encoders + linear probe).

New thesis 8.9-8.10 / paper 4.9-4.10.
"""
import sys
from pathlib import Path

from docx.shared import Inches

ROOT = Path(sys.argv[1])
sys.path.insert(0, str(ROOT / "docs" / "tools"))
from docx_tools import DocEditor  # noqa: E402

DINO = ("https://arxiv.org/abs/2304.07193",
        "M. Oquab et al., DINOv2: Learning Robust Visual Features without Supervision, Transactions on Machine "
        "Learning Research, 2024")
VIT = ("https://arxiv.org/abs/2010.11929",
       "A. Dosovitskiy et al., An Image is Worth 16x16 Words: Transformers for Image Recognition at Scale, ICLR 2021")
CURVE = ROOT / "ml" / "results" / "e2" / "e2_curve.png"

RO_E2 = [
    "Secțiunea 8.6 a ascuns o proporție aleatoare din răspunsuri. Trei verificări completează experimentul E2 "
    "(`ml/eval_missingness.py`, zece seed-uri, aceleași cazuri de test). (1) Eliminarea pe câmpuri: fiecare întrebare "
    "(grupa de vârstă, zona corpului, simptomele, textura, durata, tipul de piele) este ascunsă pentru toate cazurile, "
    "iar separat se păstrează doar acea întrebare și se ascund toate celelalte. Un câmp ascuns primește trăsături 0 și "
    "masca 1, exact ca un răspuns sărit (fără imputare, adică fără a ghici valoarea lipsă). (2) Lipsa naturală: cazurile "
    "de test sunt grupate după câte întrebări a lăsat necompletate chiar contribuitorul SCIN: complet (niciuna), parțial "
    "(unele) și gol (toate). Vârsta nu intră în numărătoare, deoarece „vârstă necunoscută” este codificată ca răspuns "
    "explicit, deci întrebarea nu lipsește niciodată. (3) Curba macro-F1 în funcție de proporția de răspunsuri ascunse "
    "(Figura 8.1). Pentru fiecare diferență se raportează media pe seed-uri, intervalul de încredere de 95% bazat pe "
    "distribuția t (pe cele zece diferențe pereche), numărul de seed-uri în care diferența este negativă și valoarea p "
    "corectată Holm (corecție pentru comparații multiple: fiecare test este judecat la un prag mai strict, astfel încât "
    "probabilitatea unui fals pozitiv în toată familia de teste să rămână 5%).",
    "Niciun câmp nu contează. Ascunderea oricărei singure întrebări schimbă macro-F1 cu cel mult 0,008, în oricare dintre "
    "cele patru modele de fuziune (concatenare fără dropout, cu dropout fix, cu dropout aleator și FiLM), iar niciun "
    "efect nu rămâne semnificativ după corecția Holm. Păstrarea unei singure întrebări dă același rezultat (scăderi de "
    "0,000-0,010, toate cu intervale care includ sau ating 0). Acest lucru confirmă secțiunile 8.5-8.7: modelul de "
    "fuziune folosește foarte puțin chestionarul, deci nu există un câmp anume a cărui lipsă să-l afecteze.",
    "Pe chestionarele lipsă în mod natural apare o diferență mică între antrenări. În medie, pe seed există 139 de cazuri "
    "de test complete, 95 parțiale și 37 goale. Pe cazurile goale, concatenarea antrenată fără răspunsuri ascunse este cu "
    "0,041 sub modelul doar cu fotografie (interval [-0,097; +0,014]), în timp ce variantele antrenate cu răspunsuri "
    "ascunse sunt la nivelul lui (+0,009 cu dropout fix, +0,006 cu dropout aleator, +0,002 pentru FiLM). Intervalele "
    "includ 0, deci este doar un indiciu: un model care nu a văzut niciodată un chestionar gol la antrenare se poate "
    "comporta mai slab când îl întâlnește. Pentru aplicație, concluzia practică rămâne aceeași: modelul doar cu "
    "fotografie este cel puțin la fel de bun la orice nivel de completare (Figura 8.1).",
]
RO_E5 = [
    "Experimentul E5 compară rețeaua EfficientNet-B0 ajustată fin (toate straturile antrenate pe SCIN, secțiunea 8.4) cu "
    "encodere înghețate urmate de o sondă liniară. Un encoder înghețat este o rețea pre-antrenată ale cărei ponderi nu se "
    "mai modifică: fiecare imagine este transformată o singură dată într-un vector de trăsături. Sonda liniară (linear "
    "probe) este un singur strat liniar antrenat peste aceste trăsături, cu aceeași pondere a claselor ca modelul de "
    "referință și cu o penalizare L2 (termen care micșorează ponderile pentru a evita supraînvățarea), aleasă pe "
    "validare. Sonda măsoară cât de utile sunt trăsăturile ca atare pentru categoriile noastre. S-au comparat trei "
    "encodere: EfficientNet-B0 cu ponderile ImageNet (aceeași rețea, dar neajustată), un Vision Transformer ViT-B/16 "
    "[{vit}] pre-antrenat supervizat pe ImageNet-21k (14 milioane de imagini, 21.000 de clase) și ViT-B/14 DINOv2 "
    "[{dino}]. Un Vision Transformer împarte imaginea în pătrate mici (patch-uri) și le prelucrează cu mecanismul de "
    "atenție, care cântărește cât contează fiecare patch pentru celelalte. DINOv2 este antrenat auto-supervizat: fără "
    "etichete, pe 142 de milioane de imagini selectate, învățând ca două vederi diferite ale aceleiași imagini să aibă "
    "trăsături asemănătoare.",
    "Rezultatele sunt în Tabelul 8.9. Sondele pe trăsături ImageNet rămân mult sub modelul ajustat fin (-0,072 și -0,069 "
    "macro-F1, mai slabe în toate cele zece seed-uri): pentru rețelele antrenate supervizat, ajustarea fină pe piele este "
    "necesară. Sonda DINOv2, în schimb, obține macro-F1 0,551 ± 0,064, cu 0,032 peste modelul ajustat fin (interval "
    "bootstrap [+0,005; +0,061], mai bună în 7 din 10 seed-uri). Diferența nu este însă confirmată la nivel de seed "
    "(t-test pereche p = 0,185, Wilcoxon p = 0,131), deoarece variația între seed-uri este mare. Acuratețea echilibrată "
    "crește de la 0,610 la 0,635, iar calibrarea este similară (ECE 0,083 față de 0,079). Pe grupuri, câștigul DINOv2 "
    "este cel mai mare pentru tonurile închise (eFST V-VI +0,108 [+0,023; +0,182]; eMST 4-6 +0,075 [+0,024; +0,122]), dar "
    "aceste diferențe nu rămân semnificative după corecția Holm; semnificative rămân doar eczema_dermatitis (+0,038, "
    "p = 0,002) și normal_other (+0,098, p = 0,027). Pentru redness_rosacea, sonda este mai slabă (-0,079).",
    "Interpretarea este că trăsăturile auto-supervizate generale se transferă la fotografiile de piele mai bine decât "
    "cele învățate pe etichetele ImageNet, încât un strat liniar peste ele egalează sau depășește o rețea mică ajustată "
    "fin pe cele 1.266 de cazuri de antrenare. Două lucruri rămân deschise. Primul: modelele specifice dermatologiei "
    "(DermLIP și PanDerm) cer aprobarea accesului la ponderi, iar setul DDI cere înregistrare; ele vor fi adăugate în "
    "același tabel. Al doilea: DINOv2 ViT-B are circa 86 de milioane de parametri, față de 5,3 milioane la "
    "EfficientNet-B0, deci este potrivit pentru inferența pe server (arhitectura aleasă pentru MVP), nu pe telefon. "
    "Alegerea modelului livrat în aplicație se va face după experimentul de calibrare și abținere (E4), comparând cele "
    "două modele pe aceleași criterii.",
]
EN_E2 = [
    "Section 4.6 hid a random share of the answers. Three checks complete experiment E2 (`ml/eval_missingness.py`, ten "
    "seeds, the same test cases). (1) Field-wise removal: each question (age group, body area, symptoms, texture, "
    "duration, skin type) is hidden for every case, and separately only that question is kept while all others are "
    "hidden. A hidden field gets features 0 and mask 1, exactly like a skipped answer (no imputation, i.e. the missing "
    "value is never guessed). (2) Natural missingness: test cases are grouped by how many questions the SCIN "
    "contributor actually left blank: complete (none), partial (some) and empty (all). Age is left out of the count "
    "because 'age unknown' is encoded as an explicit answer, so that question is never missing. (3) The curve of "
    "macro-F1 against the share of hidden answers (Figure 1). Each difference is reported as the mean over seeds, the "
    "95% confidence interval based on the t distribution (over the ten paired differences), the number of seeds where "
    "it is negative, and the Holm-adjusted p-value (a multiple-comparison correction: each test is judged at a stricter "
    "threshold so that the chance of any false positive across the family of tests stays at 5%).",
    "No field matters. Hiding any single question changes macro-F1 by at most 0.008 in any of the four fusion models "
    "(concatenation without dropout, with fixed dropout, with random dropout, and FiLM), and no effect survives the Holm "
    "correction. Keeping only one question gives the same picture (drops of 0.000-0.010, all with intervals that "
    "include or touch 0). This confirms Sections 4.5-4.7: the fusion model uses the questionnaire very little, so there "
    "is no particular field whose absence hurts it.",
    "On naturally missing questionnaires, a small difference between trainings appears. Per seed there are on average "
    "139 complete, 95 partial and 37 empty test cases. On the empty cases, concatenation trained without hidden answers "
    "is 0.041 below the photo-only model (interval [-0.097, +0.014]), while the variants trained with hidden answers are "
    "level with it (+0.009 fixed dropout, +0.006 random dropout, +0.002 FiLM). The intervals include 0, so this is only "
    "a hint: a model that never saw an empty questionnaire in training may do worse when it meets one. For the app the "
    "practical conclusion is unchanged: the photo-only model is at least as good at every level of completion "
    "(Figure 1).",
]
EN_E5 = [
    "Experiment E5 compares the fine-tuned EfficientNet-B0 (all layers trained on SCIN, Section 4.4) with frozen "
    "encoders followed by a linear probe. A frozen encoder is a pre-trained network whose weights are no longer changed: "
    "each image is turned once into a feature vector. The linear probe is a single linear layer trained on these "
    "features, with the same class weighting as the reference model and an L2 penalty (a term that shrinks the weights "
    "to avoid overfitting) chosen on validation. The probe measures how useful the features are for our categories as "
    "they are. Three encoders were compared: EfficientNet-B0 with ImageNet weights (the same network, but not "
    "fine-tuned), a Vision Transformer ViT-B/16 [{vit}] pre-trained with labels on ImageNet-21k (14 million images, "
    "21,000 classes), and DINOv2 ViT-B/14 [{dino}]. A Vision Transformer splits the image into small squares (patches) "
    "and processes them with attention, which weighs how much each patch matters to the others. DINOv2 is trained "
    "self-supervised: without labels, on 142 million curated images, by learning to give two different views of the "
    "same image similar features.",
    "Table 10 gives the results. The probes on ImageNet features stay well below the fine-tuned model (-0.072 and -0.069 "
    "macro-F1, worse in all ten seeds): for networks trained with labels, fine-tuning on skin is needed. The DINOv2 "
    "probe, in contrast, reaches macro-F1 0.551 ± 0.064, 0.032 above the fine-tuned model (bootstrap interval [+0.005, "
    "+0.061], better in 7 of 10 seeds). The difference is not confirmed at the seed level, however (paired t-test "
    "p = 0.185, Wilcoxon p = 0.131), because the spread between seeds is large. Balanced accuracy rises from 0.610 to "
    "0.635, and calibration is similar (ECE 0.083 against 0.079). By group, the DINOv2 gain is largest for darker tones "
    "(eFST V-VI +0.108 [+0.023, +0.182]; eMST 4-6 +0.075 [+0.024, +0.122]), but these differences do not survive the Holm "
    "correction; only eczema_dermatitis (+0.038, p = 0.002) and normal_other (+0.098, p = 0.027) do. For "
    "redness_rosacea the probe is weaker (-0.079).",
    "The interpretation is that general self-supervised features transfer to skin photos better than features learned "
    "from ImageNet labels, so that a linear layer on top of them matches or exceeds a small network fine-tuned on the "
    "1,266 training cases. Two points remain open. First, the dermatology-specific models (DermLIP and PanDerm) require "
    "approved access to their weights, and the DDI set requires registration; they will be added to the same table. "
    "Second, DINOv2 ViT-B has about 86 million parameters against 5.3 million for EfficientNet-B0, so it suits server "
    "inference (the architecture chosen for the MVP), not on-device use. The model shipped in the app will be chosen "
    "after the calibration and abstention experiment (E4), by comparing both models on the same criteria.",
]

E5_ROWS = [  # model, macro-F1, balanced accuracy, ECE, difference vs fine-tuned
    ["EfficientNet-B0, fine-tuned (reference)", "0.519 ± 0.040", "0.610 ± 0.038", "0.079 ± 0.017", "reference"],
    ["EfficientNet-B0 ImageNet, frozen + linear probe", "0.447 ± 0.029", "0.543 ± 0.065", "0.139 ± 0.090",
     "-0.072 [-0.099, -0.044], 0/10"],
    ["ViT-B/16 ImageNet-21k, frozen + linear probe", "0.450 ± 0.043", "0.542 ± 0.069", "0.095 ± 0.028",
     "-0.069 [-0.097, -0.039], 0/10"],
    ["ViT-B/14 DINOv2, frozen + linear probe", "0.551 ± 0.064", "0.635 ± 0.059", "0.083 ± 0.027",
     "+0.032 [+0.005, +0.061], 7/10"],
]
RO_NAMES = ["EfficientNet-B0, ajustat fin (referință)", "EfficientNet-B0 ImageNet, înghețat + sondă liniară",
            "ViT-B/16 ImageNet-21k, înghețat + sondă liniară", "ViT-B/14 DINOv2, înghețat + sondă liniară"]


def ro(s: str) -> str:
    return s.replace(", ", "; ").replace(".", ",").replace("reference", "referință")


def add_table(ed, anchor, rows, caption):
    ed.body_paragraph_before(anchor, caption)
    tbl = ed.doc.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        tbl.style = ed.doc.styles["Table Grid"]
    except KeyError:
        pass
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = v
            if i == 0:
                for r in tbl.cell(i, j).paragraphs[0].runs:
                    r.bold = True
    anchor._p.addprevious(tbl._tbl)


def add_figure(ed, anchor, path, caption):
    p = ed.body_paragraph_before(anchor, "")
    p.add_run().add_picture(str(path), width=Inches(5.8))
    ed.body_paragraph_before(anchor, caption)


DOCS = (
    ("Teza_Licenta.docx", "Capitolul 9. Concluzii și direcții viitoare",
     "8.9 Răspunsuri lipsă pe câmpuri și în mod natural (E2, completare)", RO_E2,
     "Figura 8.1. Macro-F1 pe testul SCIN în funcție de proporția de răspunsuri ascunse la testare (medie și interval "
     "de încredere de 95% pe zece seed-uri). Linia întreruptă: modelul doar cu fotografie, care nu citește chestionarul.",
     "8.10 Encodere pre-antrenate înghețate și sondă liniară (E5)", RO_E5,
     [["Model", "Macro-F1", "Acuratețe echilibrată", "ECE", "Diferență față de modelul ajustat fin"]]
     + [[n] + [ro(x) for x in r[1:]] for n, r in zip(RO_NAMES, E5_ROWS)],
     "Tabelul 8.9. Experimentul E5 pe testul SCIN (zece seed-uri; medie ± abatere standard; diferență pereche de "
     "macro-F1 cu interval bootstrap de 95% și numărul de seed-uri mai bune)."),
    ("Paper_Skin_Concern.docx", "5. Discussion",
     "4.9 Missing answers by field and natural missingness (E2, completed)", EN_E2,
     "Figure 1. SCIN test macro-F1 against the share of answers hidden at test time (mean and 95% confidence interval "
     "over ten seeds). Dashed line: the photo-only model, which does not read the questionnaire.",
     "4.10 Frozen pre-trained encoders and linear probe (E5)", EN_E5,
     [["Model", "Macro-F1", "Balanced accuracy", "ECE", "Difference vs fine-tuned model"]] + E5_ROWS,
     "Table 10. Experiment E5 on the SCIN test split (ten seeds; mean ± standard deviation; paired macro-F1 difference "
     "with 95% bootstrap interval and number of seeds better)."),
)

for path, next_head, h_e2, p_e2, fig_cap, h_e5, p_e5, rows, tab_cap in DOCS:
    ed = DocEditor(str(ROOT / "docs" / path))
    if ed.has_heading(h_e2):
        print(path, "already updated")
        continue
    n = {"vit": ed.add_reference(VIT[1], VIT[0]), "dino": ed.add_reference(DINO[1], DINO[0])}
    anchor = ed.heading(next_head)
    ed.heading_before(anchor, h_e2, 2)
    for t in p_e2:
        ed.body_paragraph_before(anchor, t.format(**n))
    add_figure(ed, anchor, CURVE, fig_cap)
    ed.heading_before(anchor, h_e5, 2)
    for t in p_e5[:1]:
        ed.body_paragraph_before(anchor, t.format(**n))
    add_table(ed, anchor, rows, tab_cap)
    for t in p_e5[1:]:
        ed.body_paragraph_before(anchor, t.format(**n))
    try:
        ed.save()
        print(path, "updated")
    except PermissionError:
        print(path, "LOCKED (open in Word) - not updated")
