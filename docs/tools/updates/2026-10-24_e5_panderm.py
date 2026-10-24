"""E5 with PanDerm-Base and PanDerm-Large: rewrite thesis 8.10 / paper 4.10 (text and table)."""
import sys
from pathlib import Path

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(sys.argv[1])
sys.path.insert(0, str(ROOT / "docs" / "tools"))
from docx_tools import DocEditor  # noqa: E402

DINO = ("https://arxiv.org/abs/2304.07193",
        "M. Oquab et al., DINOv2: Learning Robust Visual Features without Supervision, Transactions on Machine "
        "Learning Research, 2024")
VIT = ("https://arxiv.org/abs/2010.11929",
       "A. Dosovitskiy et al., An Image is Worth 16x16 Words: Transformers for Image Recognition at Scale, ICLR 2021")
PANDERM = ("https://www.nature.com/articles/s41591-025-03747-y",
           "S. Yan et al., A multimodal vision foundation model for clinical dermatology, Nature Medicine, 2025")

RO = [
    "Experimentul E5 compară rețeaua EfficientNet-B0 ajustată fin (toate straturile antrenate pe SCIN, secțiunea 8.4) cu "
    "encodere înghețate urmate de o sondă liniară. Un encoder înghețat este o rețea pre-antrenată ale cărei ponderi nu se "
    "mai modifică: fiecare imagine este transformată o singură dată într-un vector de trăsături. Sonda liniară (linear "
    "probe) este un singur strat liniar antrenat peste aceste trăsături, cu aceeași pondere a claselor ca modelul de "
    "referință și cu o penalizare L2 (termen care micșorează ponderile pentru a evita supraînvățarea), aleasă pe "
    "validare. Sonda măsoară cât de utile sunt trăsăturile ca atare pentru categoriile noastre.",
    "S-au comparat cinci encodere, în trei familii. (1) Supervizate pe ImageNet: EfficientNet-B0 (aceeași rețea, dar "
    "neajustată) și un Vision Transformer ViT-B/16 [{vit}] pre-antrenat pe ImageNet-21k (14 milioane de imagini, 21.000 "
    "de clase). Un Vision Transformer împarte imaginea în pătrate mici (patch-uri) și le prelucrează cu mecanismul de "
    "atenție, care cântărește cât contează fiecare patch pentru celelalte. (2) Auto-supervizat general: DINOv2 ViT-B/14 "
    "[{dino}], antrenat fără etichete pe 142 de milioane de imagini selectate, învățând ca două vederi diferite ale "
    "aceleiași imagini să aibă trăsături asemănătoare. (3) Model fundație pentru dermatologie: PanDerm [{panderm}], "
    "pre-antrenat auto-supervizat pe peste două milioane de imagini de boli de piele (clinice, dermatoscopice și din alte "
    "modalități), în două mărimi: Base (ViT-B/16, circa 86 de milioane de parametri) și Large (ViT-L/16, circa 303 "
    "milioane). Un model fundație este un model mare pre-antrenat pe date foarte multe și variate, gândit să fie "
    "refolosit pentru multe sarcini. Ponderile PanDerm au fost încărcate în implementarea timm a ViT; conversia a fost "
    "verificată față de codul autorilor (diferență maximă a trăsăturilor 2·10⁻⁶, adică doar rotunjire numerică).",
    "Rezultatele sunt în Tabelul 8.9. Sondele pe trăsături ImageNet rămân mult sub modelul ajustat fin (-0,072 și -0,069 "
    "macro-F1, mai slabe în toate cele zece seed-uri): pentru rețelele antrenate pe etichete ImageNet, ajustarea fină pe "
    "piele este necesară. Celelalte trei encodere depășesc modelul ajustat fin: DINOv2 cu 0,032 [+0,005; +0,061] (7/10 "
    "seed-uri, t-test pereche p = 0,185), PanDerm-Base cu 0,034 [+0,003; +0,066] (8/10, p = 0,069) și PanDerm-Large cu "
    "0,047 [+0,017; +0,078] (7/10, p = 0,022, Wilcoxon p = 0,037). Doar PanDerm-Large are un câștig confirmat și la "
    "nivel de seed. Acuratețea echilibrată crește de la 0,610 la 0,662 cu PanDerm-Large. Pe grupuri, câștigurile sunt "
    "pozitive la toate tonurile de piele (de exemplu eFST V-VI +0,084, eMST 4-6 +0,075 pentru PanDerm-Large), iar după "
    "corecția Holm rămân semnificative eMST 1-3, acne (+0,139), eczema_dermatitis și normal_other (+0,128). Pentru "
    "redness_rosacea toate sondele sunt mai slabe (PanDerm-Large -0,119), însă categoria are în medie doar 6 cazuri de "
    "test pe seed.",
    "Două limite sunt importante pentru alegerea modelului din aplicație. Prima este calibrarea: eroarea de calibrare "
    "ECE (diferența medie dintre încrederea afișată și acuratețea reală) este mai mare pentru sondele PanDerm (0,128 și "
    "0,113) decât pentru modelul ajustat fin (0,079), deci încrederea lor trebuie corectată prin scalarea temperaturii "
    "înainte de a fi arătată utilizatorului; acest lucru este testat în E4. A doua este licența: PanDerm (și DermLIP) "
    "sunt publicate sub CC BY-NC-ND 4.0, care permite doar cercetare necomercială și interzice distribuirea unor versiuni "
    "modificate, în timp ce DINOv2 are licența Apache 2.0, care permite și utilizarea comercială. Pentru lucrare, PanDerm "
    "este cel mai bun encoder; pentru o aplicație care ar putea deveni comercială, DINOv2 este varianta sigură din punct "
    "de vedere legal. Modelul DermLIP va fi adăugat în același tabel după autentificarea pe Hugging Face, iar alegerea "
    "finală se face după E4, pe aceleași criterii (macro-F1, calibrare, abținere pe tonuri de piele, licență, cost de "
    "inferență: ViT-L are de circa 57 de ori mai mulți parametri decât EfficientNet-B0).",
]
EN = [
    "Experiment E5 compares the fine-tuned EfficientNet-B0 (all layers trained on SCIN, Section 4.4) with frozen "
    "encoders followed by a linear probe. A frozen encoder is a pre-trained network whose weights are no longer changed: "
    "each image is turned once into a feature vector. The linear probe is a single linear layer trained on these "
    "features, with the same class weighting as the reference model and an L2 penalty (a term that shrinks the weights "
    "to avoid overfitting) chosen on validation. The probe measures how useful the features are for our categories as "
    "they are.",
    "Five encoders from three families were compared. (1) Supervised on ImageNet: EfficientNet-B0 (the same network, "
    "but not fine-tuned) and a Vision Transformer ViT-B/16 [{vit}] pre-trained on ImageNet-21k (14 million images, "
    "21,000 classes). A Vision Transformer splits the image into small squares (patches) and processes them with "
    "attention, which weighs how much each patch matters to the others. (2) General self-supervised: DINOv2 ViT-B/14 "
    "[{dino}], trained without labels on 142 million curated images by learning to give two different views of the "
    "same image similar features. (3) Dermatology foundation model: PanDerm [{panderm}], pre-trained self-supervised on "
    "over two million skin-disease images (clinical, dermoscopic and other modalities), in two sizes: Base (ViT-B/16, "
    "about 86 million parameters) and Large (ViT-L/16, about 303 million). A foundation model is a large model "
    "pre-trained on very much and very varied data, meant to be reused for many tasks. The PanDerm weights were loaded "
    "into the timm ViT implementation; the conversion was checked against the authors' code (maximum feature difference "
    "2e-6, i.e. numerical rounding only).",
    "Table 10 gives the results. The probes on ImageNet features stay well below the fine-tuned model (-0.072 and "
    "-0.069 macro-F1, worse in all ten seeds): for networks trained on ImageNet labels, fine-tuning on skin is needed. "
    "The other three encoders exceed the fine-tuned model: DINOv2 by 0.032 [+0.005, +0.061] (7/10 seeds, paired t-test "
    "p = 0.185), PanDerm-Base by 0.034 [+0.003, +0.066] (8/10, p = 0.069) and PanDerm-Large by 0.047 [+0.017, +0.078] "
    "(7/10, p = 0.022, Wilcoxon p = 0.037). Only PanDerm-Large has a gain confirmed at the seed level too. Balanced "
    "accuracy rises from 0.610 to 0.662 with PanDerm-Large. By group, the gains are positive for every skin tone (for "
    "example eFST V-VI +0.084 and eMST 4-6 +0.075 for PanDerm-Large), and after Holm correction eMST 1-3, acne (+0.139), "
    "eczema_dermatitis and normal_other (+0.128) remain significant. For redness_rosacea every probe is weaker "
    "(PanDerm-Large -0.119), but this category has only 6 test cases per seed on average.",
    "Two limits matter for choosing the app model. The first is calibration: the expected calibration error ECE (the "
    "mean gap between the confidence shown and the actual accuracy) is higher for the PanDerm probes (0.128 and 0.113) "
    "than for the fine-tuned model (0.079), so their confidence must be corrected with temperature scaling before it is "
    "shown to a user; this is tested in E4. The second is the licence: PanDerm (and DermLIP) are released under CC "
    "BY-NC-ND 4.0, which allows only non-commercial research and forbids distributing modified versions, whereas DINOv2 "
    "is Apache 2.0, which also allows commercial use. For the paper, PanDerm is the best encoder; for an app that might "
    "become commercial, DINOv2 is the legally safe option. DermLIP will be added to the same table after "
    "authentication on Hugging Face, and the final choice is made after E4 on the same criteria (macro-F1, "
    "calibration, abstention by skin tone, licence and inference cost: ViT-L has about 57 times more parameters than "
    "EfficientNet-B0).",
]
ROWS = [  # model, macro-F1, balanced accuracy, ECE, difference vs fine-tuned
    ["EfficientNet-B0, fine-tuned (reference)", "0.519 ± 0.040", "0.610 ± 0.038", "0.079 ± 0.017", "reference"],
    ["EfficientNet-B0 ImageNet, frozen + linear probe", "0.447 ± 0.029", "0.543 ± 0.065", "0.139 ± 0.090",
     "-0.072 [-0.099, -0.044], 0/10"],
    ["ViT-B/16 ImageNet-21k, frozen + linear probe", "0.450 ± 0.043", "0.542 ± 0.069", "0.095 ± 0.028",
     "-0.069 [-0.097, -0.039], 0/10"],
    ["ViT-B/14 DINOv2, frozen + linear probe", "0.551 ± 0.064", "0.635 ± 0.059", "0.083 ± 0.027",
     "+0.032 [+0.005, +0.061], 7/10"],
    ["PanDerm-Base ViT-B/16, frozen + linear probe", "0.553 ± 0.049", "0.642 ± 0.071", "0.128 ± 0.071",
     "+0.034 [+0.003, +0.066], 8/10"],
    ["PanDerm-Large ViT-L/16, frozen + linear probe", "0.566 ± 0.044", "0.662 ± 0.065", "0.113 ± 0.062",
     "+0.047 [+0.017, +0.078], 7/10"],
]
RO_NAMES = ["EfficientNet-B0, ajustat fin (referință)", "EfficientNet-B0 ImageNet, înghețat + sondă liniară",
            "ViT-B/16 ImageNet-21k, înghețat + sondă liniară", "ViT-B/14 DINOv2, înghețat + sondă liniară",
            "PanDerm-Base ViT-B/16, înghețat + sondă liniară", "PanDerm-Large ViT-L/16, înghețat + sondă liniară"]


def ro(s: str) -> str:
    return s.replace(", ", "; ").replace(".", ",").replace("reference", "referință")


def clear_section(ed, head_text):
    """Remove everything between a heading and the next heading; return the next heading."""
    h = ed.heading(head_text)
    nxt = h._p.getnext()
    while nxt is not None:
        if nxt.tag == qn("w:p") and Paragraph(nxt, h._parent).style.name.startswith("Heading"):
            return Paragraph(nxt, h._parent)
        gone, nxt = nxt, nxt.getnext()
        gone.getparent().remove(gone)
    raise ValueError(f"no heading after {head_text!r}")


def add_table(ed, anchor, rows, caption):
    ed.body_paragraph_before(anchor, caption)
    tbl = ed.doc.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        tbl.style = ed.doc.styles["Table Grid"]
    except KeyError:
        pass
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = v
            if i == 0:
                for r in tbl.cell(i, j).paragraphs[0].runs:
                    r.bold = True
    anchor._p.addprevious(tbl._tbl)


DOCS = (
    ("Teza_Licenta.docx", "8.10 Encodere pre-antrenate înghețate și sondă liniară (E5)", RO,
     [["Model", "Macro-F1", "Acuratețe echilibrată", "ECE", "Diferență față de modelul ajustat fin"]]
     + [[n] + [ro(x) for x in r[1:]] for n, r in zip(RO_NAMES, ROWS)],
     "Tabelul 8.9. Experimentul E5 pe testul SCIN (zece seed-uri; medie ± abatere standard; diferență pereche de "
     "macro-F1 cu interval bootstrap de 95% și numărul de seed-uri mai bune)."),
    ("Paper_Skin_Concern.docx", "4.10 Frozen pre-trained encoders and linear probe (E5)", EN,
     [["Model", "Macro-F1", "Balanced accuracy", "ECE", "Difference vs fine-tuned model"]] + ROWS,
     "Table 10. Experiment E5 on the SCIN test split (ten seeds; mean ± standard deviation; paired macro-F1 "
     "difference with 95% bootstrap interval and number of seeds better)."),
)

for path, head, paras, rows, cap in DOCS:
    ed = DocEditor(str(ROOT / "docs" / path))
    if any("PanDerm-Large" in p.text for p in ed.doc.paragraphs):
        print(path, "already updated")
        continue
    n = {"vit": ed.add_reference(VIT[1], VIT[0]), "dino": ed.add_reference(DINO[1], DINO[0]),
         "panderm": ed.add_reference(PANDERM[1], PANDERM[0])}
    anchor = clear_section(ed, head)
    for t in paras[:2]:
        ed.body_paragraph_before(anchor, t.format(**n))
    add_table(ed, anchor, rows, cap)
    for t in paras[2:]:
        ed.body_paragraph_before(anchor, t.format(**n))
    try:
        ed.save()
        print(path, "updated")
    except PermissionError:
        print(path, "LOCKED (open in Word) - not updated")
