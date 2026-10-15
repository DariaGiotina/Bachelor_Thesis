"""EER model of the app database (draw.io) and its relational mapping.

Replaces the chapter 4 placeholder of the thesis with a short chapter introduction and adds section 4.1
with Figures 4.1 (EER model) and 4.2 (relational schema). The paper covers only the ML part, so it is
not changed. Usage: python docs/tools/updates/2026-10-15_eer_model.py <repo root>
"""
import copy
import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

ROOT = Path(sys.argv[1])
sys.path.insert(0, str(ROOT / "docs" / "tools"))
from docx_tools import DocEditor  # noqa: E402

CHEN = ("https://doi.org/10.1145/320434.320440",
        "P. P. Chen, The Entity-Relationship Model: Toward a Unified View of Data, ACM Transactions on Database "
        "Systems 1(1), 9-36, 1976")
TEOREY = ("https://doi.org/10.1145/7474.7475",
          "T. J. Teorey, D. Yang, J. P. Fry, A Logical Design Methodology for Relational Databases Using the Extended "
          "Entity-Relationship Model, ACM Computing Surveys 18(2), 197-222, 1986")
DRAWIO = ("https://www.drawio.com/", "draw.io (diagrams.net), instrument de realizare a diagramelor, JGraph Ltd")

CH4 = "Capitolul 4. Proiectarea sistemului și arhitectura"
HEAD = "4.1 Modelul conceptual al datelor: diagrama Entitate-Legătură extinsă (EER)"
NEXT = "Capitolul 5. Metodologia AI/ML"

INTRO = (
    "Acest capitol descrie proiectarea aplicației: modelul conceptual al datelor, arhitectura componentelor "
    "(aplicația mobilă, API-ul, serviciul de inferență și stocarea) și interfața de programare (API). Secțiunea 4.1 "
    "prezintă modelul conceptual al datelor și transformarea lui într-o schemă relațională; arhitectura, modelul fizic "
    "în MongoDB și proiectarea API-ului vor fi descrise în secțiunile următoare."
)

RO_A = [
    "Modelul Entitate-Legătură (Entity-Relationship, ER), propus de Chen [{chen}], descrie datele unei aplicații "
    "independent de sistemul de gestiune a bazei de date care le va stoca. Elementele lui sunt: entitatea, adică o "
    "categorie de obiecte despre care se rețin date (de exemplu USER; o instanță a entității este un utilizator "
    "concret), desenată ca dreptunghi; atributul, o proprietate a entității (de exemplu email), desenat ca elipsă; "
    "cheia, atributul ale cărui valori identifică în mod unic fiecare instanță (de exemplu user_id), scrisă subliniat; "
    "și legătura (relationship), o asociere între entități (de exemplu PERFORMS: un utilizator efectuează analize), "
    "desenată ca romb. Modelul Entitate-Legătură extins (Enhanced ER, EER) adaugă concepte de tip clasă-subclasă, în "
    "special specializarea și generalizarea, cu moștenirea atributelor [{teorey}]. Diagrama a fost realizată în "
    "draw.io [{drawio}], în notația Chen, cu denumirile în limba engleză pentru a fi identice cu cele din cod. Sursa "
    "editabilă este `docs/diagrams/skincare_eer.drawio` (pagina „EER model” pentru Figura 4.1 și pagina „Relational "
    "schema” pentru Figura 4.2), iar imaginile exportate se află în același director. Fișierul este generat de "
    "scriptul `docs/tools/build_er_diagram.py`, astfel încât modificările modelului rămân reproductibile.",
    "Notația folosită are următoarele semnificații. Raportul de cardinalitate indică numărul maxim de instanțe ale "
    "unei entități care se pot asocia cu o instanță a celeilalte: 1:1 (unu la unu), 1:N (unu la mai mulți) sau M:N "
    "(mai mulți la mai mulți); valorile 1, N și M sunt scrise pe liniile legăturii. Participarea arată dacă fiecare "
    "instanță trebuie să ia parte la legătură: linia dublă marchează participarea totală (de exemplu, orice analiză "
    "aparține unui utilizator), iar linia simplă participarea parțială (un utilizator poate să nu aibă nicio "
    "analiză). O entitate slabă nu are o cheie proprie și este identificată prin entitatea de care aparține "
    "(proprietarul) împreună cu o cheie parțială; se desenează cu dreptunghi dublu, iar legătura care o identifică "
    "cu romb dublu. În model, ROUTINE_STEP este entitate slabă: pasul 3 are sens doar în interiorul unei rutine "
    "anume, deci cheia lui este (routine_id, step_no), unde step_no este cheia parțială (subliniată punctat). Un "
    "atribut compus are subatribute (full_name = first_name + last_name; skin_profile = skin_type, main_concern, "
    "climate). Un atribut multivalorat poate avea mai multe valori pentru aceeași instanță (variantele de răspuns "
    "options ale unei întrebări, funcțiile functions și marcajele de risc flags ale unui ingredient) și se desenează "
    "cu elipsă dublă. Un atribut derivat nu se stochează, ci se calculează din alte date (age_range din "
    "date_of_birth, total_duration ca sumă a duratelor pașilor unei rutine, safety_score din scorurile "
    "ingredientelor unui produs) și se desenează cu elipsă punctată.",
]

RO_B = [
    "Entitățile au fost stabilite pornind de la modelul de date din planul tezei (colecțiile users, consents, "
    "refresh_tokens, skin_analyses, products, ingredients, routines și daily_logs) și de la ecranele proiectate în "
    "Figma (chestionarul de analiză, alegerea zonelor feței, rezultatele, detaliile produsului cu scorul de siguranță, "
    "constructorul de rutine, jurnalul zilnic și setările de confidențialitate). Cele 16 entități formează patru "
    "grupuri: contul și securitatea (USER, CONSENT pentru consimțămintele GDPR, REFRESH_TOKEN pentru sesiunile de "
    "autentificare, REMINDER pentru mementourile de dimineață și seară), analiza pielii (SKIN_ANALYSIS, QUESTION, "
    "FACE_ZONE, SKIN_CONCERN), fotografiile (PHOTO cu subclasele ANALYSIS_PHOTO și PROGRESS_PHOTO) și produsele cu "
    "rutinele (PRODUCT, INGREDIENT, ROUTINE, ROUTINE_STEP, DAILY_LOG).",
    "Legăturile reflectă direct metodologia din capitolul 5. ANSWERS (M:N între SKIN_ANALYSIS și QUESTION) are "
    "atributele answer_value și is_missing: o întrebare la care utilizatorul nu a răspuns este înregistrată explicit "
    "ca lipsă, la fel ca masca de valori lipsă a modelului (secțiunea 5.5), în loc să fie completată cu o valoare "
    "inventată. PREDICTS (M:N între SKIN_ANALYSIS și SKIN_CONCERN) are atributele probability și rank, deoarece "
    "modelul întoarce o probabilitate pentru fiecare categorie din secțiunea 5.4; atributul is_uncertain al analizei "
    "marchează cazurile cu încredere scăzută, pentru care aplicația recomandă consultarea unui specialist. ASSESSES "
    "(M:N între SKIN_ANALYSIS și FACE_ZONE) reține pentru fiecare zonă a feței constatarea (finding) și gradul ei "
    "(severity); harta facială pe mai multe zone este o funcționalitate opțională (stretch), de aceea participarea "
    "este parțială, iar legătura poate rămâne goală în MVP fără ca schema să se schimbe. CONTAINS (M:N între PRODUCT "
    "și INGREDIENT) păstrează în atributul position ordinea din lista INCI (International Nomenclature of Cosmetic "
    "Ingredients, denumirile standard ale ingredientelor cosmetice), care pe etichetă este descrescătoare după "
    "concentrație. RECORDS_USE (M:N între DAILY_LOG și PRODUCT) reține perioada (AM sau PM) și momentul folosirii "
    "fiecărui produs. BASED_ON leagă opțional o rutină de analiza din care a fost generată, iar USES leagă opțional "
    "un pas de rutină de un produs concret din catalog.",
    "Componenta EER a modelului este specializarea entității PHOTO. Specializarea împarte o entitate generală "
    "(superclasa) în subclase care moștenesc toate atributele și legăturile superclasei și au, în plus, atribute și "
    "legături proprii. Specializarea este disjunctă (cercul marcat cu d): o fotografie aparține unei singure "
    "subclase. Este totală (linie dublă între PHOTO și cerc): orice fotografie este fie fotografie de analiză, fie "
    "fotografie de progres. Simbolul de submulțime de pe fiecare linie arată sensul moștenirii. Ambele tipuri au "
    "aceleași atribute de stocare (photo_id, object_key, content_type, captured_at): fotografia stă criptată în "
    "stocarea de obiecte din UE, iar baza de date reține doar cheia obiectului. Diferă însă prin rest: "
    "ANALYSIS_PHOTO are shot_type și heatmap_key (harta Grad-CAM, secțiunea 5.9) și aparține unei analize (DEPICTS, "
    "o analiză are una sau mai multe fotografii, ca în SCIN), iar PROGRESS_PHOTO are caption și documentează o zi din "
    "jurnal (DOCUMENTS, 1:1). Avantajul practic este că dreptul la ștergere din GDPR se poate aplica tuturor "
    "fotografiilor unui utilizator printr-o singură entitate.",
]

RO_C = [
    "Pentru a verifica modelul, el a fost transformat în modelul relațional (Figura 4.2), după regulile clasice de "
    "trecere de la EER la tabele [{teorey}]. În modelul relațional datele sunt păstrate în tabele (relații) formate "
    "din rânduri și coloane; cheia primară (PK) este coloana sau grupul de coloane care identifică unic fiecare rând, "
    "cheia străină (FK) este o coloană care conține cheia primară a unui rând din alt tabel, iar UQ marchează o "
    "constrângere de unicitate. Legăturile dintre tabele sunt desenate în notația „crow's foot”: liniuța înseamnă "
    "exact unu, cercul înseamnă opțional, iar „piciorul de cioară” (bifurcația) înseamnă mai mulți. Regulile aplicate "
    "au fost: fiecare entitate puternică devine un tabel, cu cheia ei drept cheie primară; atributele compuse se "
    "descompun în coloane simple; atributele derivate nu se stochează; la o legătură 1:N cheia străină se pune în "
    "tabelul de pe partea N (de exemplu SKIN_ANALYSIS.user_id); o legătură M:N devine un tabel de legătură care "
    "conține cheile ambelor entități și atributele legăturii (ANSWER, ZONE_ASSESSMENT, PREDICTION, PRODUCT_USE, "
    "PRODUCT_INGREDIENT); fiecare atribut multivalorat devine un tabel separat (QUESTION_OPTION, INGREDIENT_FUNCTION, "
    "INGREDIENT_FLAG); entitatea slabă ROUTINE_STEP primește cheia compusă (routine_id, step_no); legătura 1:1 "
    "DOCUMENTS devine o cheie străină unică (PROGRESS_PHOTO.log_id); iar specializarea lui PHOTO devine câte un tabel "
    "pentru superclasă și pentru fiecare subclasă, în care photo_id este simultan cheie primară și cheie străină, "
    "iar coloana photo_kind indică subclasa. Rezultatul are 24 de tabele.",
    "Modelul conceptual nu depinde de baza de date aleasă. Implementarea folosește MongoDB, o bază de date orientată "
    "pe documente, deci schema relațională servește ca verificare a modelului și ca reper, nu ca implementare. În "
    "MongoDB, datele citite mereu împreună cu entitatea proprietar sunt incluse (embedded) în documentul acesteia: "
    "răspunsurile, predicțiile și evaluările pe zone în documentul analizei, pașii în documentul rutinei, produsele "
    "folosite în documentul zilei din jurnal. Entitățile de referință partajate (produse, ingrediente, întrebări, "
    "categorii de probleme ale pielii) rămân colecții separate, referite prin identificator. Constrângerile de "
    "unicitate devin indexuri unice (email, barcode, inci_name și perechea user_id, log_date). Modelul fizic MongoDB, "
    "cu indexurile și regulile de validare, va fi descris în secțiunea următoare. Limitare: entitatea QUESTION "
    "urmează deocamdată întrebările SCIN mapate în `ml/features/field_mapping.md`, deoarece chestionarul final din "
    "aplicație nu este încă implementat.",
]


def add_figure(ed, anchor, image, caption, width_cm=16.0):
    par = ed.body_paragraph_before(anchor, "")
    par.paragraph_format.first_line_indent = Cm(0)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(image), width=Cm(width_cm))
    cap = anchor.insert_paragraph_before(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_toc_level2(ed, before_text, text, page):
    """Clone a level-2 table-of-contents line and insert it before the TOC line `before_text`."""
    tpl = next(p for p in ed.doc.paragraphs if p.style.name.lower() == "toc 2")
    target = ed.toc_entries(before_text)[0]
    new = copy.deepcopy(tpl._p)
    target._p.addprevious(new)
    ts = Paragraph(new, target._parent)._p.findall(".//" + qn("w:t"))
    ts[0].text, ts[-1].text = text, page
    for t in ts[1:-1]:
        t.text = "\t" if t.text == "\t" else ""


ed = DocEditor(str(ROOT / "docs" / "Teza_Licenta.docx"))
if ed.has_heading(HEAD):
    print("thesis already updated")
    sys.exit(0)
n = {"chen": ed.add_reference(CHEN[1], CHEN[0]), "teorey": ed.add_reference(TEOREY[1], TEOREY[0]),
     "drawio": ed.add_reference(DRAWIO[1], DRAWIO[0])}
ed.replace_placeholder(CH4, [INTRO])
anchor = ed.heading(NEXT)
ed.heading_before(anchor, HEAD, 2)
for t in RO_A:
    ed.body_paragraph_before(anchor, t.format(**n))
add_figure(ed, anchor, ROOT / "docs" / "diagrams" / "eer_model.png",
           "Figura 4.1. Diagrama Entitate-Legătură extinsă (EER) a aplicației, în notația Chen; componenta EER este "
           "specializarea disjunctă și totală a entității PHOTO.")
for t in RO_B:
    ed.body_paragraph_before(anchor, t.format(**n))
add_figure(ed, anchor, ROOT / "docs" / "diagrams" / "relational_schema.png",
           "Figura 4.2. Schema relațională obținută din modelul EER (24 de tabele, notația crow's foot).")
for t in RO_C:
    ed.body_paragraph_before(anchor, t.format(**n))
add_toc_level2(ed, NEXT, HEAD, "5")
try:
    ed.save()
    print("thesis updated")
except PermissionError:
    print("Teza_Licenta.docx LOCKED (open in Word) - not updated")
